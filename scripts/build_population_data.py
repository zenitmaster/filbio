#!/usr/bin/env python3
"""Build the bundled population allele-frequency JSON files.

Run once whenever a source changes; the generated JSON is committed.

    python3 -m venv .venv && .venv/bin/pip install openpyxl rdata
    curl -LO https://cran.r-project.org/src/contrib/forensicpopdata_1.0.4.tar.gz
    tar -xzf forensicpopdata_1.0.4.tar.gz
    .venv/bin/python scripts/build_population_data.py \
        --workbook "reference/<laboratory workbook>.xlsx" \
        --rda-dir forensicpopdata/data \
        --out src/data/populations

Sources
-------
* Laboratory table: sheet "Hoja1" of the laboratory's Excel workbook. The
  workbook itself is not distributed (reference/ is git-ignored): it may hold
  case data. Only the population-level frequency table is extracted.
* NIST 1036, FBI 2015 and UK DNA-17: machine-readable transcriptions shipped in
  the R package `forensicpopdata` (M. Kruijver). The underlying raw data are
  public domain (NIST, FBI) or Open Government Licence v3 (UK Home Office).

Every frequency is stored as a proportion (0-1). `chromosomes` is the number of
alleles sampled at that locus (2N); it drives the 5/(2N) minimum frequency.
"""
from __future__ import annotations

import argparse
import json
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

# --------------------------------------------------------------------------
# Laboratory workbook (Hoja1)
# --------------------------------------------------------------------------

LOCUS_ALIASES = {"THO1": "TH01"}

# The table states N = 300, but every frequency is a multiple of 1/270 and the
# observed-heterozygosity row is a multiple of 1/135, so the sample behind the
# numbers is 135 individuals (270 chromosomes).
LAB_CHROMOSOMES = 270

# 0.1 % is the workbook's placeholder for "allele not observed", not data.
LAB_FILLER_PERCENT = 0.1

# (locus, allele in workbook) -> allele it is moved to.
# D13S317 "13.2" = 4.4 %: allele 13.2 is unobserved in NIST-Hispanic, FBI-SW
# Hispanic and FBI-SE Hispanic, while allele 14 (5.7-6.1 % in all three) is
# blank in the workbook. The value sits one row above where it belongs.
LAB_CORRECTIONS = {("D13S317", "13.2"): "14"}


def canon_allele(value) -> str:
    number = float(str(value).replace("*", "").replace(",", "."))
    return str(int(number)) if number == int(number) else f"{number:g}"


def build_lab_population(workbook: Path) -> list[dict]:
    import openpyxl

    sheet = openpyxl.load_workbook(workbook, data_only=True)["Hoja1"]
    loci: dict[str, dict] = {}
    for col in range(2, 17):
        name = str(sheet.cell(1, col).value).strip()
        locus = LOCUS_ALIASES.get(name, name)
        freqs: dict[str, float] = {}
        for row in range(2, 44):
            raw = sheet.cell(row, col).value
            percent = float(str(raw).replace("*", ""))
            if abs(percent - LAB_FILLER_PERCENT) < 1e-9:
                continue
            allele = canon_allele(sheet.cell(row, 1).value)
            allele = LAB_CORRECTIONS.get((locus, allele), allele)
            freqs[allele] = round(freqs.get(allele, 0.0) + percent / 100, 6)
        ordered = dict(sorted(freqs.items(), key=lambda kv: float(kv[0])))
        loci[locus] = {"chromosomes": LAB_CHROMOSOMES, "freqs": ordered}

    return [
        {
            "id": "lab-mx-hoja1",
            "group": "lab",
            "name": {
                "es": "México — tabla del laboratorio (Excel, Hoja1)",
                "en": "Mexico — laboratory table (Excel, Hoja1)",
            },
            "individuals": LAB_CHROMOSOMES // 2,
            "source": {
                "citation": (
                    "Tabla de frecuencias del libro de Excel del laboratorio (hoja «Hoja1»). "
                    "El informe del libro declara como origen las frecuencias de población "
                    "mexicana de la SLAGF [Rev Esp Med Legal 2013;39(2):48-53]; esa "
                    "correspondencia no ha sido verificada."
                ),
                "license": "Datos del laboratorio",
            },
            "notes": {
                "es": [
                    "Tamaño muestral: la tabla indica N = 300, pero todas las frecuencias son "
                    "múltiplos de 1/270 y la fila de heterocigosidad observada lo es de 1/135. "
                    "Se usa 2N = 270 cromosomas (135 individuos).",
                    "Los valores de relleno de 0.1 % (alelos no observados) se eliminaron; en "
                    "su lugar se aplica la frecuencia alélica mínima configurada.",
                    "Corrección: D13S317 4.4 % estaba en la fila del alelo 13.2; se reasignó "
                    "al alelo 14 (13.2 no se observa en ninguna referencia hispana y el 14, "
                    "con ~6 %, estaba vacío).",
                    "Pendiente de verificar contra la fuente: D18S51 alelo 13.2 = «1.1*» "
                    "(guardado como texto) y D21S11 alelo 24.2 = 1.1 % (≈0.2 % en referencias).",
                ],
                "en": [
                    "Sample size: the table states N = 300, but every frequency is a multiple of "
                    "1/270 and the observed-heterozygosity row is a multiple of 1/135. "
                    "2N = 270 chromosomes (135 individuals) is used.",
                    "The 0.1 % filler values (unobserved alleles) were removed; the configured "
                    "minimum allele frequency is applied instead.",
                    "Correction: D13S317 4.4 % sat on the allele 13.2 row; it was reassigned to "
                    "allele 14 (13.2 is unobserved in every Hispanic reference and 14, at ~6 %, "
                    "was blank).",
                    "To verify against the source: D18S51 allele 13.2 = “1.1*” (stored as text) "
                    "and D21S11 allele 24.2 = 1.1 % (≈0.2 % in references).",
                ],
            },
            "loci": loci,
        }
    ]


# --------------------------------------------------------------------------
# forensicpopdata .rda files
# --------------------------------------------------------------------------


def _text(obj) -> str:
    value = obj.value
    return value.decode() if isinstance(value, bytes) else str(value)


def _attributes(obj) -> dict:
    out, node = {}, obj.attributes
    while node is not None and node.info.type.name == "LIST":
        tag, name = node.tag, None
        if tag is not None:
            target = tag.referenced_object if tag.info.type.name == "REF" else tag
            try:
                name = _text(target.value)
            except Exception:  # noqa: BLE001 - malformed tag, fall through
                name = None
        out[name] = node.value[0]
        node = node.value[1] if len(node.value) > 1 else None
    return out


def _names(obj) -> list[str]:
    attrs = _attributes(obj)
    if "names" in attrs:
        return [_text(x) for x in attrs["names"].value]
    for value in attrs.values():
        if value.info.type.name == "STR":
            return [_text(x) for x in value.value]
    raise ValueError("R object has no names attribute")


def load_rda(path: Path) -> dict[str, dict]:
    import rdata

    top = rdata.parser.parse_file(path).object.value[0]
    populations = {}
    for pop_name, pop in zip(_names(top), top.value):
        locus_names = _names(pop)
        n_obj = _attributes(pop).get("N")
        sizes = dict(zip(_names(n_obj), (float(v) for v in n_obj.value))) if n_obj else {}
        loci = {}
        for locus, vec in zip(locus_names, pop.value):
            freqs = {a: round(float(f), 6) for a, f in zip(_names(vec), vec.value) if f > 0}
            ordered = dict(sorted(freqs.items(), key=lambda kv: float(kv[0])))
            loci[locus] = {"chromosomes": int(round(sizes[locus])), "freqs": ordered}
        populations[pop_name] = loci
    return populations


NIST = {
    "AfAm": ("nist-afam", "Afroamericanos (EE. UU.)", "African American (U.S.)", 342),
    "Asian": ("nist-asian", "Asiáticos (EE. UU.)", "Asian (U.S.)", 97),
    "Cauc": ("nist-cauc", "Caucásicos (EE. UU.)", "Caucasian (U.S.)", 361),
    "Hisp": ("nist-hisp", "Hispanos (EE. UU.)", "Hispanic (U.S.)", 236),
}
FBI = {
    "SW Hispanic": ("fbi-sw-hispanic", "Hispanos del suroeste (EE. UU.)", "Southwestern Hispanic (U.S.)"),
    "SE Hispanic": ("fbi-se-hispanic", "Hispanos del sureste (EE. UU.)", "Southeastern Hispanic (U.S.)"),
    "Caucasian": ("fbi-caucasian", "Caucásicos (EE. UU.)", "Caucasian (U.S.)"),
    "African American": ("fbi-african-american", "Afroamericanos (EE. UU.)", "African American (U.S.)"),
    "Apache": ("fbi-apache", "Apaches", "Apache"),
    "Navajo": ("fbi-navajo", "Navajos", "Navajo"),
    "Bahamian": ("fbi-bahamian", "Bahameños", "Bahamian"),
    "Jamaican": ("fbi-jamaican", "Jamaicanos", "Jamaican"),
    "Trinidadian": ("fbi-trinidadian", "Trinitenses", "Trinidadian"),
    "Chamorro": ("fbi-chamorro", "Chamorros (Guam)", "Chamorro (Guam)"),
    "Filipino": ("fbi-filipino", "Filipinos", "Filipino"),
}
UK = {
    "White_-_EA1_&_EA2": ("uk-white", "Blancos (EA1 y EA2)", "White (EA1 & EA2)"),
    "Black_African_&_Caribbean_-_EA3": ("uk-black", "Africanos negros y caribeños (EA3)", "Black African & Caribbean (EA3)"),
    "Indian_-_EA4": ("uk-indian", "Subcontinente indio (EA4)", "Indian subcontinent (EA4)"),
    "Chinese_-_EA5": ("uk-chinese", "Chinos (EA5)", "Chinese (EA5)"),
}

NIST_SOURCE = {
    "citation": (
        "Hill CR, Duewer DL, Kline MC, Coble MD, Butler JM. U.S. population data for 29 "
        "autosomal STR loci. Forensic Sci Int Genet 2013;7:e82-e83. Revised: Steffen CR, "
        "Coble MD, Gettings KB, Vallone PM. Forensic Sci Int Genet 2017;31:e36-e40."
    ),
    "url": "https://strbase.nist.gov",
    "license": "Public domain (NIST)",
}
FBI_SOURCE = {
    "citation": (
        "Moretti TR, Moreno LI, Smerick JB, et al. Population data on the expanded CODIS core "
        "STR loci for eleven populations of significance for forensic DNA analyses in the "
        "United States. Forensic Sci Int Genet 2016;25:175-181."
    ),
    "url": "https://doi.org/10.1016/j.fsigen.2016.07.022",
    "license": "Public domain (FBI)",
}
UK_SOURCE = {
    "citation": (
        "UK Home Office. DNA population data to support the implementation of National DNA "
        "Database DNA-17 profiling."
    ),
    "url": (
        "https://www.gov.uk/government/statistics/dna-population-data-to-support-the-"
        "implementation-of-national-dna-database-dna-17-profiling"
    ),
    "license": "Open Government Licence v3.0",
}


def build_group(data: dict, table: dict, group: str, source: dict) -> list[dict]:
    out = []
    for key, meta in table.items():
        pop_id, name_es, name_en = meta[0], meta[1], meta[2]
        loci = data[key]
        individuals = meta[3] if len(meta) > 3 else max(v["chromosomes"] for v in loci.values()) // 2
        out.append(
            {
                "id": pop_id,
                "group": group,
                "name": {"es": name_es, "en": name_en},
                "individuals": individuals,
                "source": source,
                "loci": loci,
            }
        )
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--rda-dir", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    groups = {
        "lab": build_lab_population(args.workbook),
        "nist1036": build_group(load_rda(args.rda_dir / "NIST1036freqs.rda"), NIST, "nist1036", NIST_SOURCE),
        "fbi2015": build_group(load_rda(args.rda_dir / "FBI2015freqs.rda"), FBI, "fbi2015", FBI_SOURCE),
        "ukdna17": build_group(load_rda(args.rda_dir / "UKDNA17freqs.rda"), UK, "ukdna17", UK_SOURCE),
    }
    for name, populations in groups.items():
        target = args.out / f"{name}.json"
        target.write_text(json.dumps(populations, ensure_ascii=False, separators=(",", ":")) + "\n")
        n_loci = sorted({len(p["loci"]) for p in populations})
        print(f"{target}: {len(populations)} population(s), loci per population {n_loci}, {target.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
